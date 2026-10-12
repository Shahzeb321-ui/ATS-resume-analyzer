from io import BytesIO
from pathlib import Path

import streamlit as st
import torch
from docx import Document
from pypdf import PdfReader
from transformers import AutoModelForSequenceClassification, AutoTokenizer


APP_DIR = Path(__file__).resolve().parent
MODEL_DIR = APP_DIR / "resume-fit-final"
LABEL_ORDER = ("Good Fit", "Potential Fit", "No Fit")
FIT_COLORS = {
    "Good Fit": "#5a9e1b",
    "Potential Fit": "#c27a16",
    "No Fit": "#e5484d",
}

st.set_page_config(
    page_title="Resume fit predictor",
    page_icon="📄",
    layout="centered",
)

st.markdown(
    """
    <style>
    .block-container {
        max-width: 760px;
        padding-top: 2.25rem;
        padding-bottom: 3rem;
    }
    .app-subtitle {
        color: #667085;
        font-size: 0.96rem;
        margin-top: -0.5rem;
        margin-bottom: 1.5rem;
    }
    .result-card {
        min-height: 94px;
        padding: 16px;
        border-radius: 10px;
    }
    .fit-card {
        background: #eaf3df;
    }
    .confidence-card {
        background: #f5f4f1;
    }
    .result-label {
        color: #344054;
        font-size: 0.88rem;
        margin-bottom: 6px;
    }
    .result-value {
        color: #172b0b;
        font-size: 1.75rem;
        font-weight: 650;
        line-height: 1.2;
    }
    .confidence-card .result-value {
        color: #161616;
    }
    .probability-row {
        align-items: center;
        display: grid;
        gap: 12px;
        grid-template-columns: 102px minmax(50px, 1fr) 40px;
        margin: 12px 0;
    }
    .probability-label,
    .probability-value {
        color: #344054;
        font-size: 0.86rem;
    }
    .probability-value {
        text-align: right;
    }
    .probability-track {
        background: #f0efec;
        border-radius: 99px;
        height: 8px;
        overflow: hidden;
    }
    .probability-fill {
        border-radius: inherit;
        height: 100%;
    }
    div[data-testid="stFileUploader"] section {
        background: #f7f6f3;
    }
    div[data-testid="stTextArea"] textarea {
        background: #f7f6f3;
    }
    @media (max-width: 520px) {
        .block-container {
            padding-left: 1rem;
            padding-right: 1rem;
            padding-top: 1.5rem;
        }
        .probability-row {
            gap: 8px;
            grid-template-columns: 88px minmax(40px, 1fr) 36px;
        }
    }
    </style>
    """,
    unsafe_allow_html=True,
)


@st.cache_resource
def load_model(model_dir: str):
    model_path = Path(model_dir)
    if not model_path.is_dir():
        raise FileNotFoundError(f"Model directory not found: {model_path}")

    tokenizer = AutoTokenizer.from_pretrained(
        model_path,
        local_files_only=True,
    )
    model = AutoModelForSequenceClassification.from_pretrained(
        model_path,
        local_files_only=True,
    )
    model.eval()
    return tokenizer, model


def extract_resume_text(uploaded_file) -> str:
    file_bytes = uploaded_file.getvalue()
    extension = Path(uploaded_file.name).suffix.lower()

    if extension == ".pdf":
        reader = PdfReader(BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if extension == ".docx":
        document = Document(BytesIO(file_bytes))
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend(
            " ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return "\n".join(part for part in parts if part.strip()).strip()

    return file_bytes.decode("utf-8-sig").strip()


st.title("Resume fit predictor")
st.markdown(
    '<p class="app-subtitle">Upload a resume, paste a job description, and see how well they match.</p>',
    unsafe_allow_html=True,
)

uploaded_resume = st.file_uploader(
    "Upload resume",
    type=["pdf", "docx", "txt"],
    help="Supported file types: PDF, DOCX, and TXT.",
)
job_description = st.text_area(
    "Job description",
    placeholder="Paste the job description here...",
    height=112,
)

if st.button("Analyze fit", type="primary", use_container_width=True):
    if uploaded_resume is None:
        st.warning("Upload a resume before analyzing fit.")
    elif not job_description.strip():
        st.warning("Add a job description before analyzing fit.")
    else:
        resume_text = extract_resume_text(uploaded_resume)
        if not resume_text:
            st.error(
                "No readable text was found in this resume. "
                "For scanned PDFs, use a text-based PDF or upload a DOCX/TXT file."
            )
        else:
            with st.spinner("Analyzing resume and job description..."):
                tokenizer, model = load_model(str(MODEL_DIR))
                inputs = tokenizer(
                    resume_text,
                    job_description.strip(),
                    return_tensors="pt",
                    truncation=True,
                    max_length=512,
                )
                with torch.inference_mode():
                    logits = model(**inputs).logits[0]
                    probabilities = torch.softmax(logits, dim=-1).tolist()

            labels = model.config.id2label
            scores = {
                labels[index]: probability
                for index, probability in enumerate(probabilities)
            }
            predicted_label = max(scores, key=scores.get)
            st.session_state["resume_fit_result"] = (
                predicted_label,
                scores,
            )

result = st.session_state.get("resume_fit_result")
if result is not None:
    predicted_label, scores = result
    st.divider()

    fit_column, confidence_column = st.columns(2)
    fit_column.markdown(
        f"""
        <div class="result-card fit-card">
            <div class="result-label">Predicted fit</div>
            <div class="result-value">{predicted_label}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    confidence_column.markdown(
        f"""
        <div class="result-card confidence-card">
            <div class="result-label">Confidence</div>
            <div class="result-value">{scores[predicted_label]:.0%}</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.markdown("#### Class probabilities")
    for label in LABEL_ORDER:
        probability = scores[label]
        percent = probability * 100
        st.markdown(
            f"""
            <div class="probability-row">
                <span class="probability-label">{label}</span>
                <div class="probability-track">
                    <div class="probability-fill"
                         style="width: {percent:.2f}%; background: {FIT_COLORS[label]}"></div>
                </div>
                <span class="probability-value">{percent:.0f}%</span>
            </div>
            """,
            unsafe_allow_html=True,
        )
